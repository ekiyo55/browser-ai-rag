"""文書ファイルを文字にする（第12章）。

返すのは (ページ番号, 文字) の並び。ページのない形式は、ページ番号を None にした一つだけを返す。
Word の見出しは Markdown の「#」に直して、KotobaCore が見出しとして扱えるようにする。
"""

from __future__ import annotations

import re
from pathlib import Path

SUPPORTED = {".md", ".txt", ".pdf", ".docx", ".xlsx"}

Page = tuple[int | None, str]


def read_pages(path: Path) -> list[Page]:
    suffix = path.suffix.lower()
    if suffix in (".md", ".txt"):
        return [(None, path.read_text(encoding="utf-8"))]
    if suffix == ".pdf":
        return _read_pdf(path)
    if suffix == ".docx":
        return [(None, _read_docx(path))]
    if suffix == ".xlsx":
        return [(None, _read_xlsx(path))]
    raise ValueError(f"読めない形式です: {path.name}")


def _read_pdf(path: Path) -> list[Page]:
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(path)
    try:
        pages = []
        for i in range(len(pdf)):
            text = pdf[i].get_textpage().get_text_range()
            text = tidy_pdf_text(text.replace("\r\n", "\n").replace("\r", "\n"), first_page=(i == 0))
            if text.strip():
                pages.append((i + 1, text))
        return pages
    finally:
        pdf.close()


_HEADING_LIKE = re.compile(r"^(第[0-9０-９一二三四五六七八九十]+[条章節]|[0-9０-９]{1,2}[.．]\s*\S|【.+】|■)")
_SENTENCE_END = ("。", "」", "：", ":")  # 「）」で終わる行は文の途中のことが多いので含めない
_CJK = r"　-ヿ一-鿿"
_SPACE_BETWEEN = re.compile(rf"(?<=[{_CJK}]) (?=[A-Za-z0-9])|(?<=[A-Za-z0-9]) (?=[{_CJK}])")


def tidy_pdf_text(text: str, first_page: bool = False) -> str:
    """PDF から取り出した文字を整える（第12章）。

    PDF には「見出し」の印がなく、行の折り返しもそのまま改行として出てくる。
    - 1ページ目の最初の行を、文書の題名（#）にする
    - 短くて句点で終わらない行のうち、「第5条」「1. 」「【】」で始まるものを見出し（##）にする
    - 文の途中の改行（句点などで終わらない行の後ろ）をつなぐ
    - 日本語と英数字のあいだに入った余計な空白を取る（「生成 AI」→「生成AI」）
    """
    out: list[str] = []
    for line in (ln.strip() for ln in text.split("\n")):
        if not line:
            continue
        if first_page and not out:
            out.append("# " + line)
        elif len(line) <= 30 and _HEADING_LIKE.match(line) and not line.endswith("。"):
            out.append("## " + line)
        elif out and not out[-1].startswith("#") and not out[-1].endswith(_SENTENCE_END):
            out[-1] += line  # 折り返しでちぎれた行をつなぐ
        else:
            out.append(line)
    return _SPACE_BETWEEN.sub("", "\n\n".join(out))


def _read_docx(path: Path) -> str:
    import docx

    d = docx.Document(str(path))
    lines: list[str] = []
    body = d.element.body
    for block in body.iterchildren():
        tag = block.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = docx.text.paragraph.Paragraph(block, d)
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name or "") if para.style is not None else ""
            level = _heading_level(style)
            lines.append(("#" * level + " " + text) if level else text)
        elif tag == "tbl":
            table = docx.table.Table(block, d)
            rows = [[c.text.strip().replace("\n", " ") for c in r.cells] for r in table.rows]
            if rows:
                # 表は一つのかたまりで渡す。行ごとに段落を分けると、見出し行と切り離される
                table_lines = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * len(rows[0])]
                table_lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
                lines.append("\n".join(table_lines))
    return "\n\n".join(lines)


def _heading_level(style: str) -> int:
    """'Heading 2' や '見出し 2' を 2 に。'Title' / '表題' は 1。"""
    s = style.replace("見出し", "Heading").strip()
    if s in ("Title", "表題"):
        return 1
    if s.startswith("Heading"):
        tail = s[len("Heading"):].strip()
        return int(tail) if tail.isdigit() else 1
    return 0


def _read_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    out: list[str] = []
    for ws in wb.worksheets:
        rows = [["" if v is None else str(v) for v in r] for r in ws.iter_rows(values_only=True)]
        rows = [r for r in rows if any(c.strip() for c in r)]
        if not rows:
            continue
        out.append(f"## {ws.title}")
        header = rows[0]
        for r in rows[1:]:
            pairs = [f"{h}: {v}" for h, v in zip(header, r) if v.strip()]
            out.append(" / ".join(pairs))
    wb.close()
    return "\n\n".join(out)
